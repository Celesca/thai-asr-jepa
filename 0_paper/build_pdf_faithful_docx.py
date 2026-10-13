from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


HERE = Path(__file__).resolve().parent
PAGES = sorted((Path("/private/tmp/taja-docx-pages")).glob("page-*.png"))
OUTPUT = HERE / "Adapting Audio JEPA for Low Resource Thai Automatic Speech Recognition.docx"

if not PAGES:
    raise RuntimeError("No rendered PDF pages were found.")

doc = Document()
doc.add_paragraph()
section = doc.sections[0]
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)
section.header_distance = Inches(0)
section.footer_distance = Inches(0)

for index, image_path in enumerate(PAGES):
    if index:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.header_distance = Inches(0)
        section.footer_distance = Inches(0)

    paragraph = doc.paragraphs[-1] if index else doc.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Inches(0)
    paragraph.paragraph_format.space_after = Inches(0)
    paragraph.add_run().add_picture(str(image_path), width=Inches(7.5))

doc.core_properties.title = "Adapting Audio-JEPA for Low-Resource Thai Automatic Speech Recognition"
doc.core_properties.author = "Sawit Koseeyaumporn; Kittipong Tapyou; Warapol Pratumta"
doc.save(OUTPUT)
